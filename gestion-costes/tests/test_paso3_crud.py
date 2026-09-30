from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook
import io
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app import crud
from app.database import Base
from app.export import costes_a_excel, facturas_a_csv, facturas_a_excel
from app.models import EstadoPago
from app.schemas import (EmpleadoCreate, FacturaCreate, FiltroFacturas, ParteTrabajoCreate,
                         ProyectoCreate)


@pytest.fixture
def s():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as ses:
        crud.sembrar_datos_iniciales(ses)
        yield ses


def _factura(numero, fecha, nif="B12345674", nombre="Ferretería Pepe SL", **kw):
    return FacturaCreate(numero=numero, fecha_emision=fecha, proveedor={"nombre": nombre, "nif": nif},
                         lineas_iva=kw.pop("lineas", [{"tipo": 21, "base": 100}]), **kw)


def test_crud_facturas_filtros_y_resumen(s):
    obra = crud.crear_proyecto(s, ProyectoCreate(nombre="Obra Mayor", cliente="Ayto"))
    cat = crud.categoria_por_nombre(s, "material")
    crud.crear_factura(s, _factura("A-1", date(2025, 2, 10), categoria_id=cat.id, proyecto_id=obra.id))
    crud.crear_factura(s, _factura("A-2", date(2025, 5, 3), lineas=[{"tipo": 21, "base": 100}, {"tipo": 10, "base": 50}]))
    crud.crear_factura(s, _factura("X-9", date(2025, 5, 20), nif="A58818501", nombre="Gasolinera SA",
                                   estado_pago=EstadoPago.PAGADO))
    with pytest.raises(crud.Duplicado):
        crud.crear_factura(s, _factura("A-1", date(2025, 2, 10)))

    assert len(crud.listar_facturas(s, FiltroFacturas(trimestre=2, anio=2025))) == 2
    assert len(crud.listar_facturas(s, FiltroFacturas(nif="b-12345674"))) == 2
    assert [f.numero for f in crud.listar_facturas(s, FiltroFacturas(texto="gasol"))] == ["X-9"]
    # el proveedor aprende la categoría: la 2ª factura de Pepe la hereda
    assert crud.listar_facturas(s, FiltroFacturas(texto="A-2"))[0].categoria_id == cat.id

    r = crud.resumen_facturas(s, FiltroFacturas(trimestre=2, anio=2025))
    assert (r.n, r.base, r.iva, r.total) == (2, Decimal("250.00"), Decimal("47.00"), Decimal("297.00"))
    assert r.por_tipo_iva[Decimal("10")] == (Decimal("50.00"), Decimal("5.00"))
    assert r.pendiente_pago == Decimal("176.00")

    f = crud.marcar_pagada(s, crud.listar_facturas(s, FiltroFacturas(texto="A-2"))[0].id)
    assert f.fecha_pago == date.today()

    xlsx = facturas_a_excel(crud.listar_facturas(s), r, "2T 2025")
    wb = load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames == ["Facturas", "Resumen IVA"] and wb["Facturas"].max_row == 5  # cab + 3 + total
    assert facturas_a_csv(crud.listar_facturas(s)).decode("utf-8-sig").splitlines()[1].count(";") == 17


def test_partes_desglose_e_imputacion(s):
    ana = crud.crear_empleado(s, EmpleadoCreate(nombre="Ana", coste_hora=20))
    luis = crud.crear_empleado(s, EmpleadoCreate(nombre="Luis", coste_hora=15, tarifa_km=Decimal("0.30")))
    obra = crud.crear_proyecto(s, ProyectoCreate(nombre="Obra Mayor"))
    crud.crear_parte(s, ParteTrabajoCreate(empleado_id=ana.id, proyecto_id=obra.id, fecha=date(2025, 6, 2),
                                           horas=8, km=100, gastos=["Comida 20€", "Parking 5"]))
    crud.crear_parte(s, ParteTrabajoCreate(empleado_id=luis.id, proyecto_id=obra.id, fecha=date(2025, 6, 3),
                                           horas=4, km=10, comisiones=50))
    crud.crear_factura(s, _factura("B-1", date(2025, 6, 4), proyecto_id=obra.id))

    t = {c.empleado: c for c in crud.desglose_por_trabajador(s, date(2025, 6, 1), date(2025, 6, 30))}
    assert (t["Ana"].horas, t["Ana"].coste_horas, t["Ana"].importe_km, t["Ana"].dietas) == (
        Decimal("8"), Decimal("160.00"), Decimal("26.00"), Decimal("25.00"))
    assert t["Luis"].coste_total == Decimal("60.00") + Decimal("3.00") + Decimal("50.00")

    [p] = crud.coste_por_proyecto(s, date(2025, 6, 1), date(2025, 6, 30))
    assert (p.horas, p.mano_obra, p.gastos_personal, p.facturas) == (
        Decimal("12"), Decimal("220.00"), Decimal("104.00"), Decimal("100.00"))
    assert p.coste_total == Decimal("424.00")

    m = crud.metricas_dashboard(s, hoy=date(2025, 6, 15))
    assert (m.trimestre_label, m.horas, m.dietas, m.n_facturas) == ("2T 2025", Decimal("12"), Decimal("25.00"), 1)
    assert load_workbook(io.BytesIO(costes_a_excel(list(t.values()), [p]))).sheetnames == ["Por trabajador", "Por proyecto"]


def test_pago_parcial_y_ticket_sin_numero(s):
    f = crud.crear_factura(s, _factura("", date(2025, 3, 1), estado_pago="parcial", importe_pagado="50"))
    assert f.numero.startswith("SN-") and f.pendiente == Decimal("71.00")
    with pytest.raises(Exception, match="parcial"):
        _factura("Z", date(2025, 3, 1), estado_pago="parcial", importe_pagado="500")
    assert crud.resumen_facturas(s).pendiente_pago == Decimal("71.00")
    crud.actualizar_factura(s, f.id, __import__("app.schemas", fromlist=["x"]).FacturaUpdate(estado_pago="pagado"))
    assert (f.importe_pagado, f.pendiente) == (Decimal("121.00"), Decimal("0.00"))
    with pytest.raises(crud.ErrorNegocio):
        crud.actualizar_factura(s, f.id, __import__("app.schemas", fromlist=["x"]).FacturaUpdate(
            estado_pago="parcial", importe_pagado=0))


def test_excel_costes_con_detalle_partes(s):
    e = crud.crear_empleado(s, EmpleadoCreate(nombre="Ana", coste_hora=20))
    p = crud.crear_proyecto(s, ProyectoCreate(nombre="Obra"))
    crud.crear_parte(s, ParteTrabajoCreate(empleado_id=e.id, proyecto_id=p.id, fecha=date(2025, 6, 2), horas=8,
                                           gastos=["Comida 20€"]))
    wb = load_workbook(io.BytesIO(costes_a_excel([], [], "x", crud.listar_partes(s))))
    assert wb["Partes (detalle)"]["G2"].value == "Comida 20,00€"
