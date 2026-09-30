"""Exportación a Excel (formato gestoría) y CSV (separador ';', coma decimal)."""
from __future__ import annotations

import csv
import io
from collections.abc import Sequence
from datetime import date
from decimal import Decimal

from .crud import CosteProyecto, CosteTrabajador, ResumenFacturas
from .models import Factura, ParteTrabajo

FMT_EUR = '#,##0.00 "€"'
FMT_FECHA = "DD/MM/YYYY"

COLUMNAS_FACTURAS = [
    # (cabecera, ancho, función, formato)
    ("Fecha", 11, lambda f: f.fecha_emision, FMT_FECHA),
    ("Trimestre", 10, lambda f: f"{f.trimestre}T {f.anio}", None),
    ("Nº factura", 16, lambda f: f.numero, None),
    ("Proveedor", 32, lambda f: f.proveedor.nombre, None),
    ("CIF/NIF", 12, lambda f: f.proveedor.nif, None),
    ("Categoría", 20, lambda f: f.categoria.nombre if f.categoria else "", None),
    ("Cuenta", 8, lambda f: f.categoria.cuenta_contable if f.categoria else "", None),
    ("Base imponible", 14, lambda f: f.base_imponible, FMT_EUR),
    ("Tipos IVA", 12, lambda f: f.tipos_iva, None),
    ("Cuota IVA", 12, lambda f: f.cuota_iva, FMT_EUR),
    ("% IRPF", 8, lambda f: f.irpf_porcentaje or None, '0.00"%"'),
    ("Retención IRPF", 13, lambda f: f.irpf_importe, FMT_EUR),
    ("Total", 13, lambda f: f.total, FMT_EUR),
    ("Estado", 10, lambda f: f.estado_pago.value.capitalize(), None),
    ("Pagado", 12, lambda f: f.importe_pagado, FMT_EUR),
    ("Fecha pago", 11, lambda f: f.fecha_pago, FMT_FECHA),
    ("Proyecto", 20, lambda f: f.proyecto.nombre if f.proyecto else "", None),
    ("Revisar", 8, lambda f: "SÍ" if f.requiere_revision else "", None),
]
_COLS_SUMA = {"Base imponible", "Cuota IVA", "Retención IRPF", "Total"}


def _num(v):
    return float(v) if isinstance(v, Decimal) else v


def _hoja(wb, titulo: str, cabeceras: list[tuple[str, int, str | None]], filas: list[list], totales: bool = True):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(titulo)
    negrita = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="1F4E78")
    for c, (cab, ancho, _) in enumerate(cabeceras, 1):
        celda = ws.cell(row=1, column=c, value=cab)
        celda.font, celda.fill = negrita, relleno
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = ancho
    for r, fila in enumerate(filas, 2):
        for c, valor in enumerate(fila, 1):
            celda = ws.cell(row=r, column=c, value=_num(valor))
            if cabeceras[c - 1][2]:
                celda.number_format = cabeceras[c - 1][2]
    ultima = len(filas) + 1
    if totales and filas:
        fila_tot = ultima + 1
        ws.cell(row=fila_tot, column=1, value="TOTAL").font = Font(bold=True)
        for c, (_, _, fmt) in enumerate(cabeceras, 1):
            if fmt == FMT_EUR or (fmt == "0.00" and c > 1):
                letra = get_column_letter(c)
                celda = ws.cell(row=fila_tot, column=c, value=f"=SUBTOTAL(9,{letra}2:{letra}{ultima})")
                celda.number_format, celda.font = fmt, Font(bold=True)
    ws.freeze_panes = "A2"
    if filas:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cabeceras))}{ultima}"
    return ws


def facturas_a_excel(
    facturas: Sequence[Factura],
    resumen: ResumenFacturas | None = None,
    titulo_periodo: str = "",
) -> bytes:
    """Excel con hoja 'Facturas' (filtrable, con totales) y 'Resumen IVA' por tipo."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    cab = [(n, a, fmt) for n, a, _, fmt in COLUMNAS_FACTURAS]
    filas = [[fn(f) for _, _, fn, _ in COLUMNAS_FACTURAS] for f in facturas]
    _hoja(wb, "Facturas", cab, filas)

    if resumen is not None:
        ws = _hoja(
            wb,
            "Resumen IVA",
            [("Tipo IVA", 10, '0.00"%"'), ("Base imponible", 16, FMT_EUR), ("Cuota IVA", 14, FMT_EUR)],
            [[t, b, c] for t, (b, c) in resumen.por_tipo_iva.items()],
        )
        fila = ws.max_row + 2
        for etiqueta, valor in [
            ("Periodo", titulo_periodo or "-"),
            ("Nº facturas", resumen.n),
            ("Retenciones IRPF (mod. 111)", resumen.irpf),
            ("Total facturado", resumen.total),
            ("Pendiente de pago", resumen.pendiente_pago),
            ("Generado", date.today()),
        ]:
            ws.cell(row=fila, column=1, value=etiqueta).font = Font(bold=True)
            c = ws.cell(row=fila, column=2, value=_num(valor))
            c.number_format = FMT_FECHA if isinstance(valor, date) else (FMT_EUR if isinstance(valor, Decimal) else "General")
            fila += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def facturas_a_csv(facturas: Sequence[Factura]) -> bytes:
    """CSV compatible con Excel en español (';' y coma decimal, UTF-8 con BOM)."""
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow([n for n, *_ in COLUMNAS_FACTURAS])
    for f in facturas:
        fila = []
        for _, _, fn, _ in COLUMNAS_FACTURAS:
            v = fn(f)
            if isinstance(v, Decimal):
                v = f"{v:.2f}".replace(".", ",")
            elif isinstance(v, date):
                v = v.strftime("%d/%m/%Y")
            fila.append("" if v is None else v)
        w.writerow(fila)
    return buf.getvalue().encode("utf-8-sig")


def costes_a_excel(
    trabajadores: Sequence[CosteTrabajador],
    proyectos: Sequence[CosteProyecto],
    periodo: str = "",
    partes: Sequence[ParteTrabajo] = (),
) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    _hoja(
        wb,
        "Por trabajador",
        [("Empleado", 26, None), ("Partes", 8, "0"), ("Horas", 9, "0.00"), ("Coste horas", 13, FMT_EUR),
         ("Km", 9, "0.00"), ("Importe km", 12, FMT_EUR), ("Dietas", 12, FMT_EUR), ("Comisiones", 12, FMT_EUR),
         ("Total gastos", 13, FMT_EUR), ("Coste total", 13, FMT_EUR)],
        [[t.empleado, t.partes, t.horas, t.coste_horas, t.km, t.importe_km, t.dietas, t.comisiones,
          t.total_gastos, t.coste_total] for t in trabajadores],
    )
    _hoja(
        wb,
        "Por proyecto",
        [("Proyecto", 26, None), ("Cliente", 20, None), ("Horas", 9, "0.00"), ("Mano de obra", 13, FMT_EUR),
         ("Gastos personal", 14, FMT_EUR), ("Nº facturas", 10, "0"), ("Facturas (base)", 14, FMT_EUR),
         ("Coste total", 13, FMT_EUR)],
        [[p.proyecto, p.cliente or "", p.horas, p.mano_obra, p.gastos_personal, p.n_facturas, p.facturas,
          p.coste_total] for p in proyectos],
    )
    if partes:  # detalle línea a línea (útil para nóminas y justificar dietas)
        _hoja(
            wb,
            "Partes (detalle)",
            [("Fecha", 11, FMT_FECHA), ("Empleado", 22, None), ("Proyecto", 24, None), ("Horas", 8, "0.00"),
             ("Km", 8, "0.00"), ("Importe km", 11, FMT_EUR), ("Gastos", 34, None), ("Dietas", 11, FMT_EUR),
             ("Comisiones", 11, FMT_EUR), ("Total gastos", 12, FMT_EUR), ("Descripción", 40, None)],
            [[p.fecha, p.empleado.nombre, p.proyecto.nombre, p.horas, p.km, p.importe_km,
              ", ".join(f"{g.concepto} {g.importe:.2f}€".replace(".", ",") for g in p.gastos), p.importe_dietas,
              p.comisiones, p.total_gastos, p.descripcion or ""] for p in partes],
        )
    if periodo:
        for ws in wb.worksheets:
            ws.oddHeader.center.text = f"Periodo: {periodo}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
